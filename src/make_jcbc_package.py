"""
make_jcbc_package.py - assembles the upload package for J. Comput. Biophys. Chem.
(World Scientific), following its technical submission guidelines (read 2026-09-27).

usage: python src/make_jcbc_package.py
writes manuscript/journal_package/:
  Manuscript.docx                 figures embedded
  Supporting_Information.pdf      Tables S1-S5
  Monreal_Fig1.tif ...            600 dpi RGB TIFF (electronic version, in colour)
  print_CMYK/Monreal_FigN.tif     CMYK copies (the journal prints colour only from CMYK)
  Cover_Letter.docx
  CHECKLIST.md                    automatic checks + what the author must still do
"""
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from PIL import Image

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE.parent))
sys.path.insert(0, str(BASE / "src" / "manuscript"))
from check_citation_order import first_citations  # noqa: E402

SOFFICE = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")
SUB = BASE / "manuscript" / "submission"
OUT = BASE / "manuscript" / "journal_package"
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]


def main():
    from build_manuscript import KEYWORDS, RUNNING_TITLE, TITLE
    if OUT.exists():
        shutil.rmtree(OUT)
    (OUT / "print_CMYK").mkdir(parents=True)
    ms, si = SUB / "Manuscript_Audit_JCBC.docx", SUB / "Supporting_Information_Audit_JCBC.docx"
    shutil.copy(ms, OUT / "Manuscript.docx")
    shutil.copy(SUB / "Cover_Letter_JCBC.docx", OUT / "Cover_Letter.docx")
    subprocess.run([str(SOFFICE), "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(si)],
                   check=True, capture_output=True, timeout=300)
    (OUT / f"{si.stem}.pdf").rename(OUT / "Supporting_Information.pdf")

    figs = sorted((BASE / "figures").glob("Fig[0-9]*.tif"), key=lambda f: int(re.findall(r"\d+", f.stem)[0]))
    info = []
    for f in figs:
        im = Image.open(f)
        dpi = float(im.info.get("dpi", (0, 0))[0])
        info.append((f.name, dpi, im.size[0] / dpi * 25.4))
        name = f"Monreal_{f.stem}.tif"
        im.convert("RGB").save(OUT / name, dpi=(dpi, dpi), compression="tiff_lzw")
        im.convert("RGB").convert("CMYK").save(OUT / "print_CMYK" / name, dpi=(dpi, dpi), compression="tiff_lzw")

    doc = Document(ms)
    ps = [p.text for p in doc.paragraphs]
    text = "\n".join(ps)
    abstract = ps[ps.index("Abstract") + 1]
    order = first_citations(str(ms))
    caps = [p for p in ps if re.match(r"^Fig\. \d+\.", p)]
    tables = [p for p in ps if re.match(r"^Table [IVX]+ ", p)]
    body = text.split("\nReferences")[0]
    cites_roman = []
    for m in re.finditer(r"Table ([IVX]+)\b", body):
        if m.group(1) not in cites_roman and not re.search(rf"^Table {m.group(1)} ", body[m.start():m.start() + 12]):
            cites_roman.append(m.group(1))
    body_nocap = "\n".join(p for p in body.split("\n") if not re.match(r"^(Fig\. \d+\.|Table [IVX]+ )", p))
    cited_tab = []
    for n in re.findall(r"Tables? ([IVX]+)\b", body_nocap):
        if n not in cited_tab:
            cited_tab.append(n)
    si_cited = sorted(set(re.findall(r"Table S(\d+)", body)))
    refs = [p for p in ps if re.match(r"^\(\d+\) ", p)]
    title_words = set(re.split(r"[\s:,\-]+", TITLE.lower()))
    kw_clash = [kw for kw in KEYWORDS if set(re.split(r"[\s\-]+", kw.lower())) & title_words]
    widths = {n: w for n, _, w in info}
    checks = [
        ("Abstract ≤ 250 words", len(abstract.split()) <= 250, f"{len(abstract.split())} words"),
        ("3–5 keywords, none repeating a title word", 3 <= len(KEYWORDS) <= 5 and not kw_clash,
         f"{len(KEYWORDS)}; clashes: {kw_clash or 'none'}"),
        ("Running title ≤ 45 letters and spaces", len(RUNNING_TITLE) <= 45, f"{len(RUNNING_TITLE)}: {RUNNING_TITLE}"),
        ("Figures first cited in order ('Fig.')", order["Fig"] == list(range(1, len(figs) + 1)), str(order["Fig"])),
        ("'Figure' only at the start of a sentence",
         not re.search(r"[^.\n]\s\(?Figure \d", body_nocap), ""),
        ("Every figure has a caption 'Fig. N.'", len(caps) == len(figs), f"{len(caps)} captions, {len(figs)} figures"),
        ("Tables numbered with Roman numerals, cited in order", cited_tab == ROMAN[:len(tables)],
         f"cited {cited_tab}, {len(tables)} tables"),
        ("SI tables S1–S5 cited in the text", si_cited == ["1", "2", "3", "4", "5"], ", ".join(si_cited)),
        ("References: full author lists (no 'et al.')", not any("et al" in r for r in refs), f"{len(refs)} references"),
        ("Figures 600 dpi, ≤ 7 in wide", all(d >= 600 for _, d, _ in info) and all(w <= 178 for w in widths.values()),
         ", ".join(f"{n} {w:.0f} mm" for n, w in widths.items())),
        ("No leftover placeholders (the AI statement is the only one allowed)",
         sum("[AUTHOR TO COMPLETE" in p or "PENDING" in p for p in ps) <= 1,
         f"{sum('[AUTHOR TO COMPLETE' in p or 'PENDING' in p for p in ps)} placeholder(s)"),
    ]
    lines = ["# J. Comput. Biophys. Chem. submission package", "",
             "Generated by `src/make_jcbc_package.py` (guidelines read 2026-09-27).", "",
             "## Automatic checks", "", "| Check | Result | Detail |", "|---|---|---|"]
    lines += [f"| {a} | {'OK' if b else '**FAIL**'} | {c} |" for a, b, c in checks]
    lines += ["", "## Author to do before submission", "",
             ] + (["- Write the statement on the use of AI tools (yellow placeholder in the manuscript)."]
                  if any("[AUTHOR TO COMPLETE" in p for p in ps) else []) + [
              "- Reread the manuscript, the Supporting Information and the cover letter.",
              "- Optional (recommended by the journal): free Paperpal pre-submission check.",
              "- Editorial Manager: article type Research article; enter the running title and keywords.",
              "- Figures: upload the RGB TIFFs; keep the CMYK copies in print_CMYK/ for colour print "
              "(USD 80 per page, optional; the online version is in colour at no charge).",
              "- Publish the new Zenodo version (with approval) before submitting.", ""]
    (OUT / "CHECKLIST.md").write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))
    return all(b for _, b, _ in checks)


if __name__ == "__main__":
    sys.exit(0 if main() else 1)
